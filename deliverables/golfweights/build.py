"""Build the client-facing implementation workbook for golfweights.co.uk."""
import names
from names import PRODUCTS, product_h1, name_for, grams_from_slug, adapter_spec, GRAM_OVERRIDES
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HOOKS = "UK stock — same-day dispatch before 1pm and free delivery over £40."

# ---------------------------------------------------------------- helpers
BRAND_SLUG = {"TaylorMade": "taylormade", "Callaway": "callaway", "Titleist": "titleist",
              "Ping": "ping", "Scotty Cameron": "scotty-cameron", "Odyssey": "odyssey",
              "Cobra": "cobra", "PXG": "pxg", "Srixon": "srixon", "Mizuno": "mizuno",
              "Honma": "honma", "LAB Golf": "lab"}
ADAPTER_SLUG = {"Callaway": "callaway-shaft-adapters", "TaylorMade": "golf-shaft-adapters-taylormade",
                "Titleist": "titleist-shaft-adapters", "Ping": "ping-shaft-adapters",
                "Cobra": "cobra-shaft-adapters", "PXG": "pxg-shaft-adapters",
                "Srixon": "srixon-shaft-adapters", "Mizuno": "mizuno-shaft-adapters",
                "Wilson": "wilson-shaft-adapters"}
TYPE_HUB = {"driver": ("/collections/drivers", "driver head weights"),
            "putter": ("/collections/putters", "putter head weights"),
            "fairway": ("/collections/fairway-woods", "fairway wood weights"),
            "hybrid": ("/collections/hybrids", "hybrid head weights"),
            "iron": ("/collections/irons", "iron head weights")}
TYPE_WORD = {"driver": "driver", "putter": "putter", "fairway": "fairway wood",
             "hybrid": "hybrid", "iron": "iron"}
PUTTER_ONLY = {"Scotty Cameron", "Odyssey", "LAB Golf"}
# brand corrections for miscategorised rows
BRAND_FIX = {"/products/scfb": ("Scotty Cameron", "putter"),
             "/products/cszd-cszf-aftermarket-golf-weights-wrench-for-cobta-sz-speedzone-drivers-fairway-woods-hybrids-4g-14g": ("Cobra", "driver"),
             "/products/srizon-zxi-aftermarket-golf-shaft-adapter-0-335-rh-driver-fairway-woods": ("Srixon", "driver")}

KITS = ("/collections/golf-club-weight-kits-1", "golf club weight kits")
WRENCH = ("/collections/golf-club-wrench", "golf torque wrench")
ADAPT_HUB = ("/collections/shaft-adapters", "golf shaft adapters")
HOME = ("/", "golf weights")

def brand_club(p):
    if p["url"] in BRAND_FIX:
        return BRAND_FIX[p["url"]]
    club = None if p["club"] in (None, "–", "-") else p["club"]
    return p["brand"], club

def collection_anchor(brand, club):
    if brand in PUTTER_ONLY:
        return f"{brand} putter weights"
    w = TYPE_WORD.get(club or "", None)
    return f"{brand} {w} weights" if w else f"{brand} golf weights"

PRIO = {"HIGH": "High", "Med": "Medium", "Low": "Low",
        "P1": "High", "P2": "Medium", "P3": "Low"}

# ---------------------------------------------------------------- product rows
def product_row(p):
    url = p["url"]
    brand, club = brand_club(p)
    kind = p["kind"]
    h1 = product_h1(p)
    model = name_for(p)
    g = GRAM_OVERRIDES.get(url) or grams_from_slug(url)
    title = h1 if len(h1) > 55 else h1 + " | Golf Weights"
    has_brand_col = brand in BRAND_SLUG
    bcol = ("/collections/" + BRAND_SLUG[brand], collection_anchor(brand, club)) if has_brand_col else None
    hub = TYPE_HUB.get(club or "")

    # --- description
    if url == "/products/lead-tape-self-adhesive-swing-weight-tape":
        desc = "Self-adhesive golf lead tape for adding swing weight to any club — drivers, irons and putters. " + HOOKS
    elif url == "/products/sc-aftermarket-weights-for-scotty-cameron-putters":
        desc = "Replacement weights to fit all popular Scotty Cameron putter models. " + HOOKS
    elif url == "/products/pngirn-aftermarket-weights-for-ping-irons":
        desc = "Replacement head weights to fit Ping irons. " + HOOKS
    elif kind == "weights":
        opts = f"Gram options from {g}. " if g else ""
        SHORT_HOOKS = "UK stock, same-day dispatch before 1pm, free delivery over £40."
        for cand in (f"Aftermarket replacement weights to fit the {model}. {opts}" + HOOKS,
                     f"Replacement weights for the {model}. {opts}" + SHORT_HOOKS,
                     f"Replacement weights for the {model}. " + SHORT_HOOKS):
            desc = cand
            if len(cand) <= 170:
                break
    elif kind == "kit":
        SHORT_HOOKS = "UK stock, same-day dispatch before 1pm, free delivery over £40."
        for cand in (f"Complete replacement weight kit for the {model}. " + HOOKS,
                     f"Full weight kit for the {model}. " + SHORT_HOOKS):
            desc = cand
            if len(cand) <= 170:
                break
    elif kind == "adapter":
        spec = adapter_spec(url, club)
        SHORT_HOOKS = "UK stock, same-day dispatch before 1pm, free delivery over £40."
        for cand in (f"Replacement shaft adapter ({spec}) for {model} heads. " + HOOKS if spec else
                     f"Replacement shaft adapter for {model} heads. " + HOOKS,
                     f"{model} shaft adapter ({spec}). " + SHORT_HOOKS if spec else
                     f"{model} shaft adapter. " + SHORT_HOOKS):
            desc = cand
            if len(cand) <= 170:
                break
    else:  # wrench
        SHORT_HOOKS = "UK stock, same-day dispatch before 1pm, free delivery over £40."
        for cand in (f"{h1} — for fitting golf club weights, weight kits and shaft adapters. " + HOOKS,
                     f"{h1} — for fitting golf weights and adapters. " + SHORT_HOOKS):
            desc = cand
            if len(cand) <= 170:
                break

    # --- H2s
    if kind == "weights":
        h2s = ["Which Clubs These Weights Fit",
               f"Weight Options ({g})" if g else "Weight Options",
               "How to Fit Your New Weights", "Prefer a Full Kit?"]
    elif kind == "kit":
        h2s = ["What's in the Kit", "Which Clubs This Kit Fits", "How to Fit Your New Weights"]
    elif kind == "adapter":
        chart = f"Adapter Settings — See the Full {brand} Chart" if brand in ADAPTER_SLUG else "Adapter Settings"
        h2s = ["Which Clubs & Shafts This Adapter Fits", chart, "You'll Also Need a Torque Wrench"]
    else:
        h2s = ["What This Wrench Fits", "How to Use a Torque Wrench Safely"]

    # --- remove / avoid
    avoid = []
    if kind in ("weights", "kit"):
        bare = model.split(" (")[0]
        avoid.append(f"Never use the bare club name on its own as the title (e.g. just '{bare}') — always include 'Weights'.")
        if has_brand_col:
            avoid.append(f"Don't use the phrase '{collection_anchor(brand, club)}' in this page's title or headings — that phrase belongs to the {brand} collection page (link to it instead).")
    elif kind == "adapter":
        if brand in ADAPTER_SLUG:
            avoid.append(f"Don't title this page '{brand} shaft adapters' (plural) — that belongs to the {brand} adapter collection. This page is one specific adapter.")
        avoid.append("Don't add a full settings/compatibility chart here — link to the chart on the adapter collection page instead.")
    else:
        avoid.append("Don't title this page 'golf torque wrench' on its own — that phrase belongs to the Torque Wrenches collection.")
    if brand == "Odyssey" and kind in ("weights", "kit"):
        avoid.append("Three pages currently compete for 'Odyssey putter weights' — keep this page strictly about the models it fits.")
    if url == "/products/sc-aftermarket-weights-for-scotty-cameron-putters":
        avoid.append("Also avoid 'putter weights' on its own — the Putter Head Weights collection owns that phrase.")
    if url == "/products/tmst2-taylormade-stealth-2-plus-drivers-and-hybrid-aftermarket-golf-weights-2g-15g":
        avoid.append("Two of your pages currently rank for 'Stealth 2 Plus' — this one keeps it; the Stealth 2 kit page must lead with 'Weight Kit'.")
    if ("copy" in url) or url.endswith("-1") or url == "/products/scfb":
        avoid.append("This page has an odd-looking web address — keep the URL exactly as it is (do not tidy, delete or redirect it).")

    # --- links
    links = []
    if url == "/products/lead-tape-self-adhesive-swing-weight-tape":
        links = [("/collections/weights", "all golf club head weights"), HOME]
    elif kind == "weights":
        if bcol: links.append(bcol)
        if hub: links.append(hub)
        links += [KITS, WRENCH]
    elif kind == "kit":
        links.append(KITS)
        if bcol: links.append(bcol)
        if hub: links.append(hub)
        links.append(WRENCH)
    elif kind == "adapter":
        if brand in ADAPTER_SLUG:
            links.append(("/collections/" + ADAPTER_SLUG[brand], f"{brand} shaft adapters"))
        links.append(ADAPT_HUB)
        if "-lh-" in url or url.endswith("-lh"):
            links.append(("/collections/lh-shaft-adapters", "left-handed shaft adapters"))
        links.append(WRENCH)
    else:  # wrench
        links.append(("/collections/golf-club-wrench", "golf torque wrenches"))
        if bcol: links.append(bcol)
        links.append(ADAPT_HUB)

    return {"page": h1, "url": url, "prio": PRIO[p["prio"]], "title": title, "desc": desc,
            "h1": h1, "h2s": h2s, "avoid": " ".join(avoid), "links": links, "more": []}

PRODUCT_ROWS = {"weights": [], "kit": [], "adapter": [], "wrench": []}
_order = {"HIGH": 0, "Med": 1, "Low": 2}
for p in sorted(PRODUCTS, key=lambda p: (_order[p["prio"]], p["brand"])):
    PRODUCT_ROWS[p["kind"]].append(product_row(p))

# ---------------------------------------------------------------- collection rows
def C(page, url, prio, title, desc, h1, h2s, avoid, links, more=None):
    return {"page": page, "url": url, "prio": prio, "title": title, "desc": desc + " " + HOOKS,
            "h1": h1, "h2s": h2s, "avoid": avoid, "links": links, "more": more or []}

HOMEPAGE = [C("Homepage", "/", "High",
    "Golf Weights — Aftermarket Golf Club Head Weights, Kits & Adapters | UK Stock",
    "Aftermarket golf club weights, weight kits, shaft adapters and torque wrenches for every major brand.",
    "Golf Weights",
    ["Shop Golf Weights by Brand", "Shop Weights by Club Type", "Golf Club Weight Kits",
     "Golf Shaft Adapters", "Golf Torque Wrenches"],
    "Nothing to remove — but note: this is the ONLY page allowed to target 'golf weights' / 'golf club weights'. Six of your pages currently compete for that phrase; the other five are de-optimised in this sheet.",
    [TYPE_HUB["driver"], TYPE_HUB["putter"], TYPE_HUB["fairway"], TYPE_HUB["hybrid"],
     TYPE_HUB["iron"], ADAPT_HUB, KITS, ("/collections/golf-club-wrench", "golf torque wrenches")],
    [("/collections/taylormade", "TaylorMade golf weights"), ("/collections/callaway", "Callaway golf weights"),
     ("/collections/titleist", "Titleist golf weights"), ("/collections/ping", "Ping golf weights"),
     ("/collections/scotty-cameron", "Scotty Cameron putter weights"), ("/collections/odyssey", "Odyssey putter weights"),
     ("/collections/cobra", "Cobra golf weights"), ("/collections/pxg", "PXG golf weights"),
     ("/collections/srixon", "Srixon golf weights"), ("/collections/mizuno", "Mizuno putter weights"),
     ("/collections/lab", "LAB Golf putter weights"), ("/collections/honma", "Honma golf weights"),
     ("/collections/toulondesign", "Toulon Design putter weights")])]

TYPE_ROWS = [
 C("Drivers (type hub)", "/collections/drivers", "High",
   "Driver Head Weights — All Brands & Models | Golf Weights",
   "Aftermarket driver head weights — adjustable, sliding and fixed — for all major brands.",
   "Driver Head Weights",
   ["Shop Driver Weights by Brand", "Adjustable & Sliding Driver Weights", "Driver Weight Kits"],
   "Stop targeting 'drivers' on its own — this page can't outrank club retailers for it and it's the wrong buyer. Every heading pairs 'driver' with 'weights'.",
   [HOME, ("/collections/taylormade", "TaylorMade driver weights"), ("/collections/ping", "Ping driver weights"),
    ("/collections/callaway", "Callaway driver weights"), ("/collections/titleist", "Titleist driver weights"),
    ("/collections/cobra", "Cobra driver weights"), ("/collections/pxg", "PXG driver weights"), KITS],
   [("/collections/srixon", "Srixon driver weights")]),
 C("Putters (type hub)", "/collections/putters", "High",
   "Putter Head Weights — Scotty Cameron, Odyssey, Spider & More",
   "Aftermarket putter head weights and kits for Scotty Cameron, Odyssey, Spider and more.",
   "Putter Head Weights",
   ["Shop Putter Weights by Brand", "Putter Weight Kits", "Putter Weight Tools & Wrenches"],
   "Stop targeting 'putter' / 'golf putter' on their own (currently position 55 — wrong intent, unwinnable). This page — not any product page — owns 'putter weights'; the Scotty product page competing for it is being retitled in this sheet.",
   [HOME, ("/collections/scotty-cameron", "Scotty Cameron putter weights"), ("/collections/odyssey", "Odyssey putter weights"),
    ("/collections/taylormade", "TaylorMade Spider putter weights"), ("/collections/cobra", "Cobra putter weights"),
    ("/collections/mizuno", "Mizuno putter weights"), ("/collections/lab", "LAB Golf putter weights"), KITS],
   [("/collections/pxg", "PXG putter weights")]),
 C("Fairway woods (type hub)", "/collections/fairway-woods", "Medium",
   "Fairway Wood Head Weights | Golf Weights",
   "Aftermarket head weights for fairway woods — Callaway, Cobra, Ping, Titleist, TaylorMade and more.",
   "Fairway Wood Head Weights",
   ["Shop Fairway Wood Weights by Brand", "Fairway Wood Weight Kits"],
   "Stop targeting 'fairway woods' on its own (currently position 50 — people searching that want to buy a club, not a weight).",
   [HOME, ("/collections/callaway", "Callaway golf weights"), ("/collections/cobra", "Cobra golf weights"),
    ("/collections/ping", "Ping golf weights"), ("/collections/titleist", "Titleist golf weights"), KITS]),
 C("Hybrids (type hub)", "/collections/hybrids", "Medium",
   "Hybrid Head Weights | Golf Weights",
   "Aftermarket head weights for hybrids and utility clubs, all major brands.",
   "Hybrid Head Weights",
   ["Shop Hybrid Weights by Brand", "Hybrid Weight Kits"],
   "Don't target 'hybrids' or 'hybrid golf clubs' on their own — headings always pair 'hybrid' with 'weights'.",
   [HOME, ("/collections/callaway", "Callaway golf weights"), ("/collections/cobra", "Cobra golf weights"),
    ("/collections/ping", "Ping golf weights"), ("/collections/taylormade", "TaylorMade golf weights"), KITS]),
 C("Irons (type hub)", "/collections/irons", "Medium",
   "Iron Head Weights | Golf Weights",
   "Aftermarket iron head weights, including Ping iron weight replacements and PXG iron weights.",
   "Iron Head Weights",
   ["Shop Iron Weights by Brand", "Ping Iron Weight Replacements", "PXG Iron Weights"],
   "This page currently ranks for 'golf weights' (one of the six competing pages) — remove that phrase from its title and copy; keep it strictly about IRON weights.",
   [HOME, ("/collections/ping", "Ping iron weights"), ("/collections/pxg", "PXG iron weights"), KITS]),
 C("All weights (catalogue)", "/collections/weights", "Medium",
   "Shop All Golf Club Head Weights | Golf Weights",
   "Browse every aftermarket golf club head weight we stock, across all brands and club types.",
   "Shop All Golf Club Head Weights",
   ["Shop by Club Type", "Shop by Brand"],
   "Must NOT target 'golf weights' — the homepage owns that phrase and this page is one of the six currently competing for it. Keep the framing as a 'shop all' catalogue page.",
   [HOME, TYPE_HUB["driver"], TYPE_HUB["putter"], TYPE_HUB["fairway"], TYPE_HUB["hybrid"], TYPE_HUB["iron"]]),
]

BRAND_ROWS = [
 C("Scotty Cameron", "/collections/scotty-cameron", "High",
   "Scotty Cameron Putter Weights, Kits & Tools | Golf Weights",
   "Scotty Cameron putter weights, kits and tools — Newport, Phantom, Fastback and more.",
   "Scotty Cameron Putter Weights",
   ["Scotty Cameron Weights by Model (Newport, Phantom & Fastback)", "Scotty Cameron Weight Kits",
    "Scotty Cameron Weight Tools & Wrenches", "Scotty Cameron Putter Weight Chart",
    "How to Change Scotty Cameron Weights"],
   "Stop targeting 'Scotty Cameron' on its own (currently position 39 — you can't outrank Scotty and the big retailers, and those searchers want a putter, not weights). Every heading pairs the brand with 'weights'. This page already ranks #2 for 'scotty cameron putter weights' — it is the template for all the other brand pages.",
   [HOME, TYPE_HUB["putter"], KITS, ("/collections/golf-club-wrench", "golf torque wrenches"),
    ("/products/scfb", "Scotty Cameron Fastback & Squareback weights"),
    ("/products/scotty-cameron-red-insert-phantom-x-newport-series-golo-series-futura-series-phantom-series-california-series-button-series-special-select-my-girl-putter-aftermarket-golf-weights-5g-45g", "Scotty Cameron Newport & Phantom weights")]),
 C("TaylorMade", "/collections/taylormade", "High",
   "TaylorMade Golf Weights & Kits | Driver & Spider Putter Weights",
   "TaylorMade driver and Spider putter weights — Qi10, Qi35, Stealth, SIM and BRNR.",
   "TaylorMade Golf Weights",
   ["TaylorMade Driver Weights", "TaylorMade Spider & TP Putter Weights",
    "TaylorMade Fairway & Hybrid Weights", "TaylorMade Weight Kits", "TaylorMade Shaft Adapters"],
   "Remove bare model names as targets (e.g. 'TaylorMade Qi10' alone) — model phrases belong on product pages, always with 'weights' added. This page currently competes for 'golf weights' — keep it strictly TaylorMade.",
   [HOME, TYPE_HUB["driver"], TYPE_HUB["putter"],
    ("/collections/golf-shaft-adapters-taylormade", "TaylorMade shaft adapters"), KITS,
    ("/products/taylormade-brnr-mini-driver-1-2-aftermarket-golf-club-weights-2g-17g", "TaylorMade BRNR Mini driver weights"),
    ("/products/taylormade-2016-m2-driver-aftermarket-golf-weights", "TaylorMade M2 driver weights"),
    ("/products/taylormade-qi-10-aftermarket-max-ls-hl-golf-club-driver-fairway-hybrid-weights-2-17g", "TaylorMade Qi10 driver weights")]),
 C("Odyssey", "/collections/odyssey", "High",
   "Odyssey Putter Weights & Weight Kits | Ai-ONE, Jailbird & White Hot",
   "Odyssey putter weights and kits — Ai-ONE, Jailbird, White Hot and Stroke Lab.",
   "Odyssey Putter Weights",
   ["Odyssey Ai-ONE Weights", "Odyssey Stroke Lab & White Hot Weights",
    "Odyssey Jailbird & Toulon Weights", "Odyssey Weight Kits"],
   "This page ALONE owns 'Odyssey putter weights' — two product pages currently compete with it for that phrase and are being retitled to model-only names in this sheet.",
   [HOME, TYPE_HUB["putter"], KITS, ("/collections/golf-club-wrench", "golf torque wrenches"),
    ("/products/odyssey-ai-one-ai-crusier-series-putter-aftermarket-golf-weights-5g-35g", "Odyssey Ai-ONE putter weights"),
    ("/products/odyssey-stroke-lab-series-white-hot-og-series-ten-series-o-works-putters-aftermarket-golf-weights-kit", "Odyssey Stroke Lab weight kit")]),
 C("PXG", "/collections/pxg", "High",
   "PXG Golf Weights & Weight Kits | Driver, Iron & Putter",
   "PXG weights and weight kits — Black Ops drivers, Gen-series irons, Battle Ready putters.",
   "PXG Golf Weights",
   ["PXG Driver Weights (Black Ops & Gen Series)", "PXG Iron Weights",
    "PXG Putter Weights (Battle Ready)", "PXG Weight Kits", "PXG Weight Chart"],
   "Stop targeting 'PXG' on its own (currently position 19 — unwinnable, wrong intent). Every heading pairs PXG with 'weights'.",
   [HOME, TYPE_HUB["driver"], TYPE_HUB["iron"], ("/collections/pxg-shaft-adapters", "PXG shaft adapters"), KITS,
    ("/products/pxg-black-ops-gen6-0311-gen5-0311-gen4-0811-0317-0211-drivers-fairway-woods-hybrids-aftermarket-golf-weights-kit", "PXG Black Ops weight kit"),
    ("/products/pxg-gen-8-irons-aftermarket-golf-weights", "PXG Gen 8 iron weights")],
   [(TYPE_HUB["putter"])]),
 C("Callaway", "/collections/callaway", "Medium",
   "Callaway Golf Weights & Kits | Driver & Fairway Wood Weights",
   "Aftermarket Callaway driver and fairway wood weights — Elyte, Paradym, Rogue, Mavrik, Epic and Apex.",
   "Callaway Golf Weights",
   ["Callaway Driver Weights", "Callaway Fairway Wood Weights",
    "Callaway Apex Iron & Hybrid Weights", "Callaway Weight Kits", "Callaway Shaft Adapters"],
   "Don't target bare model names ('Callaway Paradym' alone) — model phrases live on product pages with 'weights' added.",
   [HOME, TYPE_HUB["driver"], TYPE_HUB["fairway"],
    ("/collections/callaway-shaft-adapters", "Callaway shaft adapters"), KITS,
    ("/products/callaway-paradym-driver-and-fairway-wood-aftermarket-golf-weights-4g-18g", "Callaway Paradym driver weights"),
    ("/products/callaway-elyte-x-max-fast-driver-fairway-wood-aftermarket-golf-weights", "Callaway Elyte driver weights")]),
 C("Ping", "/collections/ping", "Medium",
   "Ping Golf Weights & Kits | Driver & Iron Weights (G425–G440)",
   "Ping driver weights, iron weight replacements and kits — G425, G430, G440 and more.",
   "Ping Golf Weights",
   ["Ping Driver Weights (G30–G440)", "Ping Fairway & Hybrid Weights",
    "Ping Iron Weight Replacements", "Ping Putter Weights (Vault)", "Ping Weight Kits"],
   "Keep the iron weight replacement content — this page already ranks #1 for 'ping iron weight replacement'. Don't target bare model names ('Ping G425 driver' alone).",
   [HOME, TYPE_HUB["driver"], TYPE_HUB["iron"], ("/collections/ping-shaft-adapters", "Ping shaft adapters"), KITS,
    ("/products/quality-aftermarket-sliding-head-weight-for-ping-g425-driver", "Ping G425 driver weights"),
    ("/products/ping-g430-driver-aftermarket-golf-weights-5g-34g", "Ping G430 driver weights"),
    ("/products/ping-g440-driver-aftermarket-golf-weights", "Ping G440 driver weights")]),
 C("Titleist", "/collections/titleist", "Medium",
   "Titleist Golf Weights & Kits | SureFit Driver & Fairway Weights",
   "Aftermarket Titleist SureFit weights and weight kits for GT, TSR, TS and 913–917 drivers and fairways.",
   "Titleist Golf Weights",
   ["Titleist Driver Weights (GT, TSR & TS)", "Titleist Fairway & Hybrid Weights",
    "Titleist SureFit Weight Kits", "Titleist Shaft Adapters & SureFit Chart"],
   "This page currently competes for 'golf weights' — keep it strictly Titleist. Weight-settings searches are big for Titleist: cross-link the adapter page's SureFit chart prominently.",
   [HOME, TYPE_HUB["driver"], ("/collections/titleist-shaft-adapters", "Titleist shaft adapters"), KITS,
    ("/products/titleistts3-driver-aftermarket-magnetic-surefit-golf-weights-7g-19g", "Titleist TS3 driver weights"),
    ("/products/titleist-gt3-driver-fairway-wood-hybrid-aftermarket-golf-weights", "Titleist GT3 driver weights")]),
 C("Cobra", "/collections/cobra", "Medium",
   "Cobra Golf Weights | Driver, Fairway & Putter Weights",
   "Aftermarket Cobra weights — Darkspeed, DS-Adapt, Aerojet, LTDx, Radspeed drivers and King putters.",
   "Cobra Golf Weights",
   ["Cobra Driver Weights (Darkspeed, Aerojet, LTDx & Radspeed)", "Cobra Fairway & Hybrid Weights",
    "Cobra King Putter Weights", "Cobra Weight Kits", "Cobra Shaft Adapters"],
   "Don't target bare model names ('Cobra Aerojet' alone — its product page currently ranks #12 for that, wrong intent; it's being retitled to '...Weights').",
   [HOME, TYPE_HUB["driver"], TYPE_HUB["putter"], ("/collections/cobra-shaft-adapters", "Cobra shaft adapters"), KITS,
    ("/products/cobra-aerojet-max-ls-golf-club-driver-fairway-woods-hybrids-aftermarket-golf-weights-4g-16g", "Cobra Aerojet driver weights"),
    ("/products/cobra-king-vintage-series-putter-aftermarket-golf-weights-5g-25g", "Cobra King putter weights")]),
 C("Srixon", "/collections/srixon", "Low",
   "Srixon Golf Weights | ZXi & ZX Driver Weights",
   "Aftermarket Srixon driver weights and kits for the ZXi and ZX ranges.",
   "Srixon Golf Weights",
   ["Srixon Driver Weights (ZXi & ZX)", "Srixon Weight Kits", "Srixon Shaft Adapters"],
   "Keep this page light — Srixon's real search demand is shaft adapters, so the adapter link should be prominent (banner-level).",
   [HOME, TYPE_HUB["driver"], ("/collections/srixon-shaft-adapters", "Srixon shaft adapters"), KITS,
    ("/products/srixon-zxi-zxi-ls-zxi-max-driver-aftermarket-golf-weights", "Srixon ZXi driver weights")]),
 C("Mizuno", "/collections/mizuno", "Low",
   "Mizuno Putter Weights | M-Craft Weights & Kits",
   "Aftermarket Mizuno putter weights for the M-Craft range, plus ST-series driver weights.",
   "Mizuno Putter Weights",
   ["Mizuno M-Craft Putter Weights", "Mizuno ST Driver Weights", "Mizuno Weight Kits", "Mizuno Shaft Adapters"],
   "This page currently competes for 'golf weights' — keep it strictly Mizuno.",
   [HOME, TYPE_HUB["putter"], ("/collections/mizuno-shaft-adapters", "Mizuno shaft adapters"), KITS,
    ("/products/mz-aftermarket-weights-for-mizuno-st-range-drivers", "Mizuno ST driver weights")]),
 C("LAB Golf", "/collections/lab", "Low",
   "LAB Golf Putter Weights | DF & Mezz",
   "Aftermarket putter weights for LAB Golf putters, including the DF and Mezz ranges.",
   "LAB Golf Putter Weights",
   ["LAB Golf Weights by Model", "How LAB Putter Weighting Works"],
   "Light page — rising brand. Keep headings paired with 'weights'.",
   [HOME, TYPE_HUB["putter"],
    ("/products/lab-golf-club-putter-head-aftermarket-weights", "LAB Golf putter head weights")]),
 C("Honma", "/collections/honma", "Low",
   "Honma Golf Weights | T//World Driver Weights",
   "Aftermarket weights for Honma T//World drivers and fairway woods.",
   "Honma Golf Weights",
   ["Honma T//World Weights"],
   "Minimal effort page — negligible search demand. Just apply the title/H1 and move on.",
   [HOME, TYPE_HUB["driver"],
    ("/products/hnm-aftermarket-weights-for-honma-golf-t-world-tr20-gs-driver-fairway-woods", "Honma T//World driver weights")]),
 C("Toulon Design", "/collections/toulondesign", "Low",
   "Toulon Design Putter Weights | Golf Weights",
   "Aftermarket putter weights for Toulon Design putters, 2022 and 2025 collections.",
   "Toulon Design Putter Weights",
   ["Toulon Weights by Collection"],
   "Minimal effort page — negligible search demand. Just apply the title/H1 and move on.",
   [HOME, TYPE_HUB["putter"],
    ("/products/odyssey-toulon-2025-series-golf-club-putter-head-aftermarket-weights", "Toulon 2025 putter weights"),
    ("/products/odyssey-toulon-2022-series-golf-club-putter-head-aftermarket-weights", "Toulon 2022 putter weights")]),
]

ADAPTER_ROWS = [
 C("Shaft adapters (hub)", "/collections/shaft-adapters", "High",
   "Golf Shaft Adapters — All Brands | Sleeves & Adapter Charts",
   "Golf shaft adapters and sleeves for every major brand, with adapter settings and compatibility charts.",
   "Golf Shaft Adapters",
   ["Shop Shaft Adapters by Brand", "Adapter Settings & Compatibility Charts",
    "Left-Handed Shaft Adapters", "Torque Wrenches for Fitting"],
   "—",
   [HOME, ("/collections/titleist-shaft-adapters", "Titleist shaft adapters"),
    ("/collections/golf-shaft-adapters-taylormade", "TaylorMade shaft adapters"),
    ("/collections/callaway-shaft-adapters", "Callaway shaft adapters"),
    ("/collections/ping-shaft-adapters", "Ping shaft adapters"),
    ("/collections/lh-shaft-adapters", "left-handed shaft adapters"),
    ("/collections/golf-club-wrench", "golf torque wrenches")],
   [("/collections/cobra-shaft-adapters", "Cobra shaft adapters"),
    ("/collections/pxg-shaft-adapters", "PXG shaft adapters"),
    ("/collections/srixon-shaft-adapters", "Srixon shaft adapters"),
    ("/collections/mizuno-shaft-adapters", "Mizuno shaft adapters"),
    ("/collections/wilson-shaft-adapters", "Wilson shaft adapters")]),
 C("Titleist shaft adapters", "/collections/titleist-shaft-adapters", "High",
   "Titleist Shaft Adapters + SureFit Loft Chart | Golf Weights",
   "Titleist SureFit shaft adapters with the full SureFit settings chart (A•1 to D•4).",
   "Titleist Shaft Adapters & SureFit Chart",
   ["Titleist SureFit Adapter Chart (A•1–D•4 Settings)", "Titleist Driver Adapters",
    "Titleist Fairway & Hybrid Adapters", "Which Adapter Fits Your Titleist?",
    "You'll Need a Torque Wrench"],
   "BIGGEST single opportunity in this whole plan: put the full SureFit settings/compatibility chart ON this page (as a table, not an image). Searches for the chart alone are ~1,400/month in the US and 100/month in the UK with zero competition difficulty.",
   [ADAPT_HUB, ("/collections/titleist", "Titleist golf weights"),
    ("/collections/golf-club-wrench", "golf torque wrenches"),
    ("/products/titleist-aftermarket-surefit-golf-sleeve-adapter-0-335-rh-lh-driver", "Titleist SureFit driver adapter")]),
 C("TaylorMade shaft adapters", "/collections/golf-shaft-adapters-taylormade", "High",
   "TaylorMade Shaft Adapters + Loft Sleeve Chart | Golf Weights",
   "TaylorMade loft sleeve adapters with the full chart — which sleeve fits M, SIM, Stealth and Qi.",
   "TaylorMade Shaft Adapters & Loft Sleeve Chart",
   ["TaylorMade Loft Sleeve Compatibility Chart (M, SIM, Stealth & Qi)",
    "TaylorMade Driver Adapters", "TaylorMade Fairway & Hybrid Adapters",
    "TaylorMade Adapter Settings Explained"],
   "Your adapter product already ranks #3 for 'compatibility chart' searches — move the chart UP to this collection page and have products link to it. Do NOT change this page's unusual web address.",
   [ADAPT_HUB, ("/collections/taylormade", "TaylorMade golf weights"),
    ("/collections/golf-club-wrench", "golf torque wrenches"),
    ("/products/taylormade-aftermarket-golf-sleeve-adapter-0-335-rh-driver-fairway", "TaylorMade .335 RH shaft adapter")]),
 C("Callaway shaft adapters", "/collections/callaway-shaft-adapters", "High",
   "Callaway Shaft Adapters + OptiFit Settings Chart | Golf Weights",
   "Callaway OptiFit shaft adapters for drivers and fairways, with the full OptiFit settings chart.",
   "Callaway Shaft Adapters & OptiFit Chart",
   ["Callaway OptiFit Settings Chart", "Callaway Driver Adapters",
    "Callaway Fairway & Hybrid Adapters", "Which Adapter Fits Your Callaway?"],
   "Add the OptiFit settings chart on-page — 'callaway adapter settings' is ~500/month in the US.",
   [ADAPT_HUB, ("/collections/callaway", "Callaway golf weights"),
    ("/collections/golf-club-wrench", "golf torque wrenches"),
    ("/products/callaway-elyte-paradym-rogue-epic-aftermarket-golf-shaft-adapter-0-335-lh-driver", "Callaway Elyte & Paradym LH adapter")]),
 C("Ping shaft adapters", "/collections/ping-shaft-adapters", "High",
   "Ping Shaft Adapters (Gen 1–3) + Settings Chart | Golf Weights",
   "Ping shaft adapters across generations 1–3, with a compatibility table covering G400 to G440.",
   "Ping Shaft Adapters & Settings Chart",
   ["Ping Adapter Compatibility by Generation (G400–G440)", "Ping Driver Adapters",
    "Ping Fairway & Hybrid Adapters", "Ping Adapter Settings Chart"],
   "Add the generation compatibility table (which adapter fits G400/G410/G425/G430/G440).",
   [ADAPT_HUB, ("/collections/ping", "Ping golf weights"),
    ("/collections/golf-club-wrench", "golf torque wrenches"),
    ("/products/ping-g430-g425-g410-aftermarket-golf-sleeve-adapter-0-335-rh-driver-fairway-woods", "Ping G430/G425/G410 adapter")]),
 C("Cobra shaft adapters", "/collections/cobra-shaft-adapters", "Medium",
   "Cobra Shaft Adapters + Loft Chart | Golf Weights",
   "Cobra shaft adapters for Darkspeed, DS-Adapt, Aerojet, LTDx and Radspeed, with loft settings chart.",
   "Cobra Shaft Adapters & Loft Chart",
   ["Cobra Adapter Loft Chart", "Cobra Driver Adapters", "Cobra Fairway Adapters"],
   "—",
   [ADAPT_HUB, ("/collections/cobra", "Cobra golf weights"),
    ("/collections/golf-club-wrench", "golf torque wrenches")]),
 C("PXG shaft adapters", "/collections/pxg-shaft-adapters", "Medium",
   "PXG Shaft Adapters + Settings Chart | Golf Weights",
   "PXG shaft adapters for Black Ops and Gen-series drivers, fairways and hybrids, with settings chart.",
   "PXG Shaft Adapters & Settings Chart",
   ["PXG Adapter Settings Chart", "PXG Driver Adapters", "PXG Hybrid Adapters"],
   "—",
   [ADAPT_HUB, ("/collections/pxg", "PXG golf weights"),
    ("/collections/golf-club-wrench", "golf torque wrenches")]),
 C("Srixon shaft adapters", "/collections/srixon-shaft-adapters", "Medium",
   "Srixon Shaft Adapters + Settings | Golf Weights",
   "Srixon shaft adapters for the ZXi, ZX and Z-Series ranges, with adapter settings.",
   "Srixon Shaft Adapters",
   ["Srixon Adapter Settings", "Srixon Driver Adapters"],
   "Srixon's best search cluster is adapters (not weights) — this page leads the Srixon effort.",
   [ADAPT_HUB, ("/collections/srixon", "Srixon golf weights"),
    ("/collections/golf-club-wrench", "golf torque wrenches")]),
 C("Mizuno shaft adapters", "/collections/mizuno-shaft-adapters", "Low",
   "Mizuno Shaft Adapters | Golf Weights",
   "Mizuno shaft adapters for the ST range of drivers.",
   "Mizuno Shaft Adapters", ["Mizuno Driver Adapters"],
   "Minimal effort — apply the title/H1 and move on.",
   [ADAPT_HUB, ("/collections/mizuno", "Mizuno putter weights")]),
 C("Wilson shaft adapters", "/collections/wilson-shaft-adapters", "Low",
   "Wilson Shaft Adapters | Golf Weights",
   "Wilson shaft adapters, including the Dynapower series.",
   "Wilson Shaft Adapters", ["Wilson Driver Adapters"],
   "Minimal effort — apply the title/H1 and move on.",
   [ADAPT_HUB]),
 C("Left-handed shaft adapters", "/collections/lh-shaft-adapters", "Medium",
   "Left-Handed Golf Shaft Adapters | Golf Weights",
   "Left-handed golf shaft adapters across TaylorMade, Callaway, Cobra, Ping, PXG and Wilson.",
   "Left-Handed Shaft Adapters",
   ["LH Adapters by Brand", "Adapter Settings & Charts"],
   "—",
   [ADAPT_HUB,
    ("/products/callaway-elyte-paradym-rogue-epic-aftermarket-golf-shaft-adapter-0-335-lh-driver", "Callaway LH driver adapter"),
    ("/products/ping-g440-aftermarket-golf-shaft-adapter-0-335-lh-driver-fairway-woods", "Ping G440 LH adapter"),
    ("/products/cobra-ds-adapt-aftermarket-golf-shaft-adapter-0-335-lh-driver", "Cobra DS-Adapt LH adapter"),
    ("/products/taylormade-qi35-stealth-sim-aftermarket-golf-shaft-adapter-0-335-lh-driver-fairway-woods", "TaylorMade Qi35 LH adapter")]),
]

KIT_WRENCH_ROWS = [
 C("Weight kits (hub)", "/collections/golf-club-weight-kits-1", "High",
   "Golf Club Weight Kits | TaylorMade, PXG, Titleist & More",
   "Complete golf club weight kits by brand — TaylorMade, PXG, Titleist, Scotty Cameron, Odyssey and more.",
   "Golf Club Weight Kits",
   ["TaylorMade Weight Kits", "PXG Weight Kits", "Titleist Weight Kits",
    "Scotty Cameron Weight Kits", "Odyssey Weight Kits", "More Brands"],
   "Do NOT change this page's web address (the '-1' on the end must stay). Per-brand kit searches are strong — the brand H2 sections above each need their own product carousel.",
   [HOME, ("/collections/taylormade", "TaylorMade golf weights"), ("/collections/pxg", "PXG golf weights"),
    ("/collections/titleist", "Titleist golf weights"), ("/collections/scotty-cameron", "Scotty Cameron putter weights"),
    ("/collections/golf-club-wrench", "golf torque wrenches"),
    ("/products/taylormade-qi-10-aftermarket-max-ls-hl-golf-club-driver-fairway-hybrid-weights-kit", "TaylorMade Qi10 weight kit"),
    ("/products/pxg-black-ops-gen6-0311-gen5-0311-gen4-0811-0317-0211-drivers-fairway-woods-hybrids-aftermarket-golf-weights-kit", "PXG Black Ops weight kit")]),
 C("Torque wrenches (hub)", "/collections/golf-club-wrench", "High",
   "Golf Torque Wrenches (T25 & T20) | Callaway, TaylorMade, Ping & More",
   "Golf torque wrenches — universal T25 and T20 wrenches plus putter weight tools for every major brand.",
   "Golf Torque Wrenches",
   ["Universal T25 Torque Wrench", "T20 & T15 Wrenches",
    "Putter Weight Tools (Scotty Cameron & Odyssey)", "Which Wrench Do You Need?"],
   "Your T25 universal wrench already ranks #3–4 for 'golf torque wrench' searches — this collection consolidates that cluster. Every weight and adapter page links here ('you'll need a wrench').",
   [HOME, ADAPT_HUB, KITS, ("/collections/scotty-cameron", "Scotty Cameron putter weights"),
    ("/products/t25-universal-torque-wrench-for-taylormade-titleist-callaway-ping-cobra-odyssey-pxg-mizuno-honma", "T25 universal torque wrench"),
    ("/products/t20-torque-wrench-for-taylormade-callaway-pxg", "T20 torque wrench")]),
]

MERCH_ROWS = [
 C(name, url, "Low", f"{plain} | Golf Weights",
   f"{plain} at Golf Weights — aftermarket golf club weights, kits and adapters.",
   plain, [],
   "Keep this page plain — NO keyword-targeted copy or headings. It must not compete with the SEO pages.",
   [])
 for name, url, plain in [
    ("Best sellers", "/collections/best-sellers", "Best Sellers"),
    ("New products", "/collections/new-products", "New Products"),
    ("Sale", "/collections/sale", "Sale"),
 ]]

SECTIONS = [
 ("HOMEPAGE", HOMEPAGE),
 ("CLUB-TYPE COLLECTIONS", TYPE_ROWS),
 ("BRAND COLLECTIONS — GOLF WEIGHTS", BRAND_ROWS),
 ("SHAFT ADAPTER COLLECTIONS", ADAPTER_ROWS),
 ("WEIGHT KITS & WRENCHES COLLECTIONS", KIT_WRENCH_ROWS),
 ("MERCHANDISING COLLECTIONS (no SEO targeting — leave plain)", MERCH_ROWS),
 ("PRODUCTS — WEIGHTS", PRODUCT_ROWS["weights"]),
 ("PRODUCTS — WEIGHT KITS", PRODUCT_ROWS["kit"]),
 ("PRODUCTS — SHAFT ADAPTERS", PRODUCT_ROWS["adapter"]),
 ("PRODUCTS — WRENCHES & TOOLS", PRODUCT_ROWS["wrench"]),
]

# ---------------------------------------------------------------- sanity checks
all_rows = [r for _, rows in SECTIONS for r in rows]
titles = [r["title"] for r in all_rows]
from collections import Counter
dupes = [t for t, c in Counter(titles).items() if c > 1]
assert not dupes, f"Duplicate titles: {dupes}"
urls = [r["url"] for r in all_rows]
dupes = [u for u, c in Counter(urls).items() if c > 1]
assert not dupes, f"Duplicate URLs: {dupes}"
print(f"{len(all_rows)} rows, all titles & URLs unique")

# ---------------------------------------------------------------- workbook
wb = Workbook()
FONT = "Arial"
GREEN = "1F4D3A"
LIGHT = "E7F0EA"
GREY = "F5F5F5"
thin = Side(style="thin", color="D0D0D0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def style_header(ws, ncols, height=32):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=GREEN)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = height

def section_row(ws, row, text, ncols):
    ws.cell(row=row, column=1, value=text)
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, size=10, bold=True, color=GREEN)
        cell.fill = PatternFill("solid", fgColor=LIGHT)
    ws.row_dimensions[row].height = 20

def body_cell(cell, bold=False):
    cell.font = Font(name=FONT, size=10, bold=bold)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    cell.border = BORDER

# ============ Tab 1: Read me ============
ws = wb.active
ws.title = "Read Me First"
ws.sheet_view.showGridLines = False
readme = [
 ("Golf Weights — SEO Implementation Sheet", ""),
 ("", ""),
 ("What this is", "Every commercial page on golfweights.co.uk with its new page title, meta description, H1 heading and suggested H2 section headings — all copy-and-paste ready. The 'Internal Links' tab lists the links to add on each page."),
 ("", ""),
 ("THE GOLDEN RULES", ""),
 ("1. Never change a web address", "Page URLs stay exactly as they are — even odd-looking ones ending in '-copy', '-1' or short codes like 'scfb'. Changing a URL loses its rankings. This sheet only changes titles, descriptions, headings, text and links."),
 ("2. One search phrase per page", "Each page targets ONE phrase of its own. If a phrase appears in another page's title, keep it off yours — the 'Remove / avoid on this page' column tells you exactly what to keep off each page. This fixes the current problem where six of your pages compete with each other for 'golf weights'."),
 ("3. No two pages share a title", "Every title in this sheet is unique. When you edit or add pages later, keep it that way — near-identical titles make Google treat pages as interchangeable and they cannibalise each other."),
 ("4. Copy text exactly", "Titles, descriptions, headings and link anchor text are written to be pasted as-is."),
 ("5. Work High → Medium → Low", "The Priority column tells you the order. 'High' pages are the ones already ranking on page 1–2 of Google — they will move first."),
 ("", ""),
 ("HOW TO APPLY IN SHOPIFY", ""),
 ("Collections", "Shopify admin → Products → Collections → open the collection. Set Title to the 'New H1' value. Then scroll to 'Search engine listing' → Edit, and paste the 'New Page Title' and 'New Meta Description'."),
 ("Products", "Shopify admin → Products → open the product. Set the product Title to the 'New H1' value. Then in 'Search engine listing' → Edit, paste the 'New Page Title' and 'New Meta Description'."),
 ("H2 headings", "The H2 suggestions match the new page template: each H2 introduces a section/product carousel on the page (e.g. 'TaylorMade Driver Weights' above the driver products). Add them in the order given. On product pages they are ordinary content headings (fitment list, weight options, how-to)."),
 ("Internal links", "On the 'Internal Links' tab, each row lists the links to add ON that page. 'Links to' is the destination; 'Anchor text' is the exact clickable wording. Add them naturally in the page's intro or description text."),
 ("", ""),
 ("Delivery claims used", "Descriptions use: UK stock · same-day dispatch before 1pm · free delivery over £40. If any of these change, let us know and we'll update the wording."),
 ("Questions?", "Anything unclear — ask before publishing. It's much easier to fix in this sheet than after Google has re-crawled the site."),
]
for i, (a, b) in enumerate(readme, start=1):
    ca = ws.cell(row=i, column=1, value=a or None)
    cb = ws.cell(row=i, column=2, value=b or None)
    ca.font = Font(name=FONT, size=14 if i == 1 else 10, bold=True,
                   color=GREEN if (a and not b) else "000000")
    cb.font = Font(name=FONT, size=10)
    cb.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 30
ws.column_dimensions["B"].width = 110

# ============ Tab 2: Page edits ============
ws = wb.create_sheet("Page Edits")
headers = ["Page", "Page URL (do not change)", "Priority", "New Page Title",
           "New Meta Description", "New H1"] + [f"H2 suggestion {i}" for i in range(1, 7)] + \
          ["Remove / avoid on this page"]
ws.append(headers)
style_header(ws, len(headers))
r = 2
for sec, rows in SECTIONS:
    section_row(ws, r, sec, len(headers)); r += 1
    for row in rows:
        h2s = (row["h2s"] + [""] * 6)[:6]
        vals = [row["page"], row["url"], row["prio"], row["title"], row["desc"], row["h1"]] + h2s + [row["avoid"]]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c, value=v or None)
            body_cell(cell, bold=(c == 1))
        if row["prio"] == "High":
            ws.cell(row=r, column=3).font = Font(name=FONT, size=10, bold=True, color="B00000")
        r += 1
widths = [34, 46, 9, 52, 58, 38, 28, 28, 28, 28, 28, 28, 55]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "C2"

# ============ Tab 3: Internal links ============
ws = wb.create_sheet("Internal Links")
NL = 8
headers = ["Page", "Page URL"]
for i in range(1, NL + 1):
    headers += [f"Link {i} — links to", f"Link {i} — anchor text"]
headers += ["More links to add (one per line: URL — anchor text)"]
ws.append(headers)
style_header(ws, len(headers))
r = 2
for sec, rows in SECTIONS:
    section_row(ws, r, sec, len(headers)); r += 1
    for row in rows:
        links = row["links"][:NL]
        extra = row["links"][NL:] + row["more"]
        vals = [row["page"], row["url"]]
        for i in range(NL):
            if i < len(links):
                vals += [links[i][0], links[i][1]]
            else:
                vals += ["", ""]
        vals.append("\n".join(f"{u} — {a}" for u, a in extra))
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c, value=v or None)
            body_cell(cell, bold=(c == 1))
        r += 1
ws.column_dimensions["A"].width = 34
ws.column_dimensions["B"].width = 46
for i in range(3, 3 + NL * 2):
    ws.column_dimensions[get_column_letter(i)].width = 34 if i % 2 == 1 else 30
ws.column_dimensions[get_column_letter(3 + NL * 2)].width = 55
ws.freeze_panes = "C2"

out = "GolfWeights_SEO_Implementation_Sheet.xlsx"
wb.save(out)
print("saved", out)
